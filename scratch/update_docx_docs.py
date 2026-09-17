"""
Update Weekly_Time_Management_Plan_UPDATED.docx in C:\\Knowledge-Base\\docs.
"""

import sys
from pathlib import Path
import docx

docx_path = r"C:\Knowledge-Base\docs\Weekly_Time_Management_Plan_UPDATED.docx"

print(f"Loading {docx_path}...")
doc = docx.Document(docx_path)

# Update paragraphs if any
for p in doc.paragraphs:
    if "Capstone" in p.text:
        p.text = p.text.replace("Capstone", "BAT-600")

# Update table cells
for table in doc.tables:
    for row_idx, row in enumerate(table.rows):
        for col_idx, cell in enumerate(row.cells):
            text = cell.text.strip()
            # Monday / Wednesday 1:00 PM - 2:00 PM
            if "ONLINE CLASS" in text and "1:00" in text:
                cell.text = "BAT-600\nONLINE CLASS\n1:00–2:00"
            # Tuesday 9:00 AM - 10:00 AM
            elif "ONLINE CLASS" in text and "9:00" in text:
                cell.text = "ICC-600\nONLINE CLASS\n9:00–10:00"

doc.save(docx_path)
print(f"Successfully saved updated docx: {docx_path}")

# Verify
doc2 = docx.Document(docx_path)
print("\n--- Verified Table 0 Cells ---")
for r in doc2.tables[0].rows:
    row_text = [c.text.replace('\n', ' ').strip() for c in r.cells]
    print(row_text)
